#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-properties-json.py  --  rebuild public/properties.json (the viewer's data file)

The file is derived from two things:
  * the survey FIELDS   : the Gravity Forms export (form 2)
  * the survey FILES    : (pic / vid) the first image / first video of each survey

Modes
  --prove      Rebuild the July 2026 file from source/merged_properties_final.csv and compare it with an existing
               properties.json cell by cell (per-key match percentage, every mismatch class). Nothing is written.
  (default)    Build from the portal dump (wp_gf_entry.csv + wp_gf_entry_meta.csv) and historical-surveys.xlsx:
               the survey set and the files come from the workbook (01_Surveys / 04_Attachments), the fields from the
               dump. Writes --out.

Usage (from this folder)
  python build-properties-json.py --prove --previous ../../public/properties.json
  python build-properties-json.py --dump-dir "<...>/oct-surveys-dump" --workbook historical-surveys.xlsx \
          --previous ../../public/properties.json --out ../../public/properties.json \
          --changes-csv <somewhere>/properties-changes.csv

Rules (all verified by --prove: 100% of the 7,398 July records are reproduced except the documented precision case)
  * text values are stripped (also non-breaking spaces); pandas' default "missing" tokens (N/A, NA, NULL, None, null,
    nan, n/a, ...) become "".
  * dc: date_created, but the zero date 0000-00-00 00:00:00 becomes "".
  * am, nf: numbers written as floats ("290000" -> "290000.0", "3" -> "3.0").
  * lat, lng: kept as separate text fields (Gravity Forms 124.2 / 124.1), never JSON.
  * pic: first picture; vid: first video (see pick_media). In the legacy CSV mode that is the first URL of the
    Pictures list (or the first non-video URL of Backup_Documents), and the first URL with a video extension of
    Videos + Backup_Documents. In the default mode they come from 04_Attachments (exact originalUrl text); a missing
    file, a placeholder (.../2020/12/1) or a URL without a file extension gives "".
  * _s: lower-cased search text: the non-empty values of id on cp cn bd ad ci al ai ar ak m4 m7 ms pt ps ta af am dl
    no eb sb ob po joined with a single space.
  * output: JSON array, 4-space indent, UTF-8 (not escaped), CRLF line ends, no trailing newline (as the old file).
"""
import argparse, collections, csv, json, os, re, sys, urllib.parse

csv.field_size_limit(2 ** 31 - 1)

# viewer key -> (column label of merged_properties_final.csv, Gravity Forms meta key)
# NB: in the CSV the labels of the last two area columns are swapped against the form: the column "Areas_In_Karachi"
# holds field 117 (text "Market/Area") and "Market_Area_117" holds field 143 (select "Areas In Karachi"). The viewer's
# ak / m7 follow the CSV columns, so ak = 117 and m7 = 143.
FIELDS = collections.OrderedDict([
    ('id', ('id', None)), ('dc', ('date_created', None)), ('st', ('status', None)),
    ('on', ('Owner_Name', '68')), ('cp', ('Contact_Person', '69')), ('cn', ('Contact_Number', '70')),
    ('bd', ('Business_Developer', '110')), ('ad', ('Address', '111')), ('ci', ('City', '112')),
    ('al', ('Areas_In_Lahore', '113')), ('ai', ('Areas_In_Islamabad', '115')), ('ar', ('Areas_In_Rawalpindi', '116')),
    ('ak', ('Areas_In_Karachi', '117')), ('m4', ('Market_Area_114', '114')), ('m7', ('Market_Area_117', '143')),
    ('lng', ('Longitude', '124.1')), ('lat', ('Latitude', '124.2')), ('ms', ('Map_Search_Address', '124.3')),
    ('pt', ('Property_Type', '136')), ('ps', ('Plot_Size', '137')), ('pa', ('Plot_Area', '138')),
    ('cs', ('Covered_Size', '142')), ('sf', ('Select_Floors', '141')), ('ta', ('Total_Area', '119')),
    ('af', ('Available_For', '134')), ('rp', ('Rent_Sale_Price_Available', '132')),
    ('nt', ('No_Rent_Price_Text', '135')), ('am', ('Amount', '76')), ('dl', ('Deal', '74')),
    ('av', ('Advance', '75')), ('sc', ('Security', '77')), ('po', ('Possession', '78')),
    ('gp', ('Grace_Period', '126')), ('inc', ('Increment', '127')), ('ap', ('Agreement_Period', '128')),
    ('no', ('Note', '122')), ('eb', ('Earlier_Brands', '120')), ('sb', ('Select_Brands', '125')),
    ('ob', ('Other_Brands', '118')), ('an', ('Available_Not_Rented', '133')), ('nf', ('No_of_Floors', '3')),
    ('f1t', ('Floor_1_Type', '26')), ('f1a', ('Floor_1_Area', '4')), ('f1s', ('Floor_1_Size', '47')),
    ('f2t', ('Floor_2_Type', '84')), ('f2a', ('Floor_2_Area', '5')), ('f2s', ('Floor_2_Size', '48')),
    ('f3t', ('Floor_3_Type', '85')), ('f3a', ('Floor_3_Area', '6')), ('f3s', ('Floor_3_Size', '49')),
])
# media columns (not part of FIELDS): Pictures = 83, Videos = 103, Backup_Documents = 129
MEDIA = {'pics': ('Pictures', '83'), 'vids': ('Videos', '103'), 'backup': ('Backup_Documents', '129')}
KEYS = list(FIELDS) + ['pic', 'vid', '_s']
SEARCH_KEYS = ['id', 'on', 'cp', 'cn', 'bd', 'ad', 'ci', 'al', 'ai', 'ar', 'ak', 'm4', 'm7', 'ms', 'pt', 'ps', 'ta',
               'af', 'am', 'dl', 'no', 'eb', 'sb', 'ob', 'po']
FLOAT_KEYS = ('am', 'nf')
PANDAS_NA = {'', '#N/A', '#N/A N/A', '#NA', '-1.#IND', '-1.#QNAN', '-NaN', '-nan', '1.#IND', '1.#QNAN', '<NA>', 'N/A',
             'NA', 'NULL', 'NaN', 'None', 'n/a', 'nan', 'null'}
VIDEO_EXT = {'.mp4', '.mov', '.avi', '.3gp', '.mkv', '.webm'}
IMAGE_EXT = {'.jpg', '.jpeg', '.png'}


# ------------------------------------------------------------------------------------------------------------------
# readers: each returns {id: {viewer key: text, 'pics': [..], 'vids': [..], 'backup': [..]}}
# ------------------------------------------------------------------------------------------------------------------
def clean(v):
    v = (v or '').strip()
    return '' if v in PANDAS_NA else v


def url_list(v):
    try:
        l = json.loads(v or '')
    except ValueError:
        return []
    return [u.strip() for u in l if isinstance(u, str) and u.strip()] if isinstance(l, list) else []


def read_merged_csv(path):
    out = {}
    with open(path, newline='', encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            rec = {k: clean(r.get(col)) for k, (col, _) in FIELDS.items()}
            rec['id'] = r['id']
            rec['dc'] = r['date_created']
            rec['st'] = r['status']
            for name, (col, _) in MEDIA.items():
                rec[name] = url_list(r.get(col))
            out[int(r['id'])] = rec
    return out


def read_dump(dump_dir):
    """entry + meta CSV of the portal dump. A meta key saved more than once keeps every DISTINCT non-empty value,
    sorted, joined with " | " (this is how merged_properties_final.csv was made; verified with the July dump)."""
    entries, metas = {}, collections.defaultdict(lambda: collections.defaultdict(set))
    with open(os.path.join(dump_dir, 'wp_gf_entry.csv'), newline='', encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            if r['form_id'] == '2':
                entries[int(r['id'])] = r
    with open(os.path.join(dump_dir, 'wp_gf_entry_meta.csv'), newline='', encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            i = int(r['entry_id'])
            if i in entries and r['meta_value'] not in (None, 'NULL', ''):
                metas[i][r['meta_key']].add(r['meta_value'].strip())
    out = {}
    for i, e in entries.items():
        m = {k: ' | '.join(sorted(v for v in vs if v)) for k, vs in metas.get(i, {}).items()}
        rec = {}
        for k, (_, mk) in FIELDS.items():
            rec[k] = clean(m.get(mk)) if mk else ''
        rec['id'] = str(i)
        rec['dc'] = e['date_created']
        rec['st'] = e['status']
        for name, (_, mk) in MEDIA.items():
            rec[name] = [u for v in sorted(metas.get(i, {}).get(mk, ())) for u in url_list(v)]
        out[i] = rec
    return out


def read_workbook_media(path):
    """-> (survey ids in sheet order, {id: {'IMAGE': [urls], 'VIDEO': [urls]}}) from 01_Surveys / 04_Attachments"""
    try:
        import openpyxl
    except ImportError:
        sys.exit('openpyxl is required:  pip install openpyxl')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ids = [int(r[0]) for r in wb['01_Surveys'].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    rows = wb['04_Attachments'].iter_rows(values_only=True)
    hdr = list(next(rows))
    ci = {h: hdr.index(h) for h in ('legacySurveyId', 'attachmentType', 'originalUrl')}
    media = collections.defaultdict(lambda: {'IMAGE': [], 'VIDEO': []})
    for r in rows:
        if r[ci['legacySurveyId']] is not None:
            media[int(r[ci['legacySurveyId']])][r[ci['attachmentType']]].append(r[ci['originalUrl']])
    wb.close()
    return ids, media


# ------------------------------------------------------------------------------------------------------------------
# record building
# ------------------------------------------------------------------------------------------------------------------
def ext(u):
    return os.path.splitext(urllib.parse.urlparse(u).path)[1].lower()


def usable(u):
    """a real file: has a file extension (the placeholder .../2020/12/1 has none)"""
    return bool(u) and bool(ext(u))


def pick_media_legacy(rec):
    """how the July file picked pic / vid from the three upload lists"""
    pics, vids, backup = rec['pics'], rec['vids'], rec['backup']
    pic = (pics or [u for u in backup if ext(u) not in VIDEO_EXT][:1] or [''])[0]
    vid = next((u for u in vids + backup if ext(u) in VIDEO_EXT), '')
    return pic, vid


def pick_media_workbook(media, sid):
    m = media.get(sid, {'IMAGE': [], 'VIDEO': []})
    pic = next((u for u in m['IMAGE'] if usable(u)), '')
    vid = next((u for u in m['VIDEO'] if usable(u)), '')
    return pic, vid


def fmt_float(v):
    if not v:
        return ''
    try:
        return repr(float(v.replace(',', '')))
    except ValueError:
        return v


def make_record(rec, pic, vid):
    r = collections.OrderedDict()
    for k in FIELDS:
        v = rec[k]
        if k == 'dc' and v.startswith('0000'):
            v = ''
        elif k in FLOAT_KEYS:
            v = fmt_float(v)
        r[k] = v
    r['pic'], r['vid'] = pic, vid
    r['_s'] = ' '.join(str(r[k]) for k in SEARCH_KEYS if str(r[k]).strip()).lower()
    return r


def dumps(records):
    return json.dumps(records, indent=4, ensure_ascii=False).replace('\n', '\r\n')


def write_json(path, records):
    with open(path, 'wb') as f:
        f.write(dumps(records).encode('utf-8'))


# ------------------------------------------------------------------------------------------------------------------
# modes
# ------------------------------------------------------------------------------------------------------------------
def same_number(a, b, tol=1e-9):
    try:
        return abs(float(a) - float(b)) <= tol
    except ValueError:
        return False


def prove(a):
    src = read_merged_csv(a.merged_csv)
    old = json.load(open(a.previous, encoding='utf-8'))
    order_ok = [int(x['id']) for x in old] == list(src)
    built = [make_record(src[i], *pick_media_legacy(src[i])) for i in src]
    print('records: built %d, existing %d; same ids in the same order: %s; same key order: %s'
          % (len(built), len(old), order_ok, all(list(b) == list(o) for b, o in zip(built, old))))
    byid = {x['id']: x for x in old}
    match, total = collections.Counter(), 0
    mism = collections.defaultdict(list)
    for b in built:
        o = byid.get(b['id'])
        if o is None:
            continue
        total += 1
        for k in KEYS:
            if b[k] == o[k]:
                match[k] += 1
            else:
                mism[k].append((b['id'], o[k], b[k]))
    print('\nper-key match (%d records)' % total)
    for k in KEYS:
        print('  %-4s %7.3f%%  (%d differ)' % (k, 100.0 * match[k] / total, total - match[k]))
    for k, v in mism.items():
        mx = max(abs(float(o) - float(n)) for _, o, n in v) if k in ('lat', 'lng') else None
        print('\n%s: %d mismatches%s; first: %s' % (k, len(v), '; max |diff| %.2e degrees' % mx if mx else '', v[:3]))
    same_bytes = dumps(old).encode('utf-8') == open(a.previous, 'rb').read()
    print('\nserialising the EXISTING file with the writer reproduces it byte for byte: %s' % same_bytes)
    print('cells total %d, identical %d (%.4f%%)' % (total * len(KEYS), sum(match.values()),
                                                      100.0 * sum(match.values()) / (total * len(KEYS))))


def build(a):
    ids, media = read_workbook_media(a.workbook)
    src = read_dump(a.dump_dir)
    missing = [i for i in ids if i not in src]
    if missing:
        sys.exit('survey ids in the workbook but not in the dump: %s' % missing[:20])
    old = {x['id']: x for x in json.load(open(a.previous, encoding='utf-8'))} if a.previous else {}
    out, changes = [], []
    for i in sorted(ids):
        pic, vid = pick_media_workbook(media, i)
        r = make_record(src[i], pic, vid)
        o = old.get(str(i))
        if o is None:
            changes.append((i, '*', '', 'NEW SURVEY'))
        else:
            for k in ('lat', 'lng'):           # old file: float text parsed by pandas' fast parser (last digit off)
                if o[k] and r[k] and o[k] != r[k] and same_number(o[k], r[k]):
                    r[k] = o[k]
            for k in KEYS:
                if k != '_s' and r[k] != o[k]:
                    changes.append((i, k, o[k], r[k]))
        out.append(r)
    if a.changes_csv:
        with open(a.changes_csv, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(['id', 'field', 'old', 'new'])
            w.writerows(changes)
    write_json(a.out, out)
    new = [c for c in changes if c[1] == '*']
    print('wrote %s: %d records (%d new, %d existing; %d changed cells in existing records)'
          % (a.out, len(out), len(new), len(out) - len(new), len(changes) - len(new)))
    print('changes by field:', dict(collections.Counter(c[1] for c in changes if c[1] != '*')))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--prove', action='store_true', help='rebuild the July file from the CSV and compare (no writes)')
    ap.add_argument('--merged-csv', default='source/merged_properties_final.csv')
    ap.add_argument('--dump-dir', help='folder with wp_gf_entry.csv and wp_gf_entry_meta.csv')
    ap.add_argument('--workbook', default='historical-surveys.xlsx')
    ap.add_argument('--previous', help='existing properties.json (change report; keeps lat/lng text equal within 1e-9)')
    ap.add_argument('--out', default='../../public/properties.json')
    ap.add_argument('--changes-csv', help='write every changed cell (id, field, old, new) to this CSV')
    a = ap.parse_args()
    if a.prove:
        if not a.previous:
            ap.error('--prove needs --previous')
        prove(a)
    else:
        if not a.dump_dir:
            ap.error('give --dump-dir')
        build(a)


if __name__ == '__main__':
    main()
