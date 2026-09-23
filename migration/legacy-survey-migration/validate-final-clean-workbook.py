#!/usr/bin/env python3
"""Validate legacy-survey-final-clean-7398.xlsx against the spec and the source CSV.

Usage (from the repo root):

    python migration/legacy-survey-migration/validate-final-clean-workbook.py

Writes final-clean-validation-report.json and exits 1 on any error.
The expected column lists are written out here on purpose (not imported from
the generator) so a generator mistake cannot validate itself.
"""
from __future__ import annotations

import sys

sys.dont_write_bytecode = True

import csv  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
from collections import Counter, defaultdict  # noqa: E402
from datetime import date, datetime  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import normalize as N  # noqa: E402

SOURCE = HERE / "source" / "merged_properties_final.csv"
HEAD = HERE / "source" / "attachment_head_scan.csv"
XLSX = HERE / "legacy-survey-final-clean-7398.xlsx"
REPORT = HERE / "final-clean-validation-report.json"

SHEETS = ["01_Surveys", "02_Survey_Contacts", "03_Attachments", "04_Exceptions"]
SURVEY_FIELDS = """serialNumber propertyName address locationCoordinates city area zipCode
yearBuilt yearRenovated buildingClass propertyType ownerName ownershipEntity propertyManager
acquisitionDate plotArea areaUnit plotSize landShape topography zoning grossBuildingArea
rentableArea usableArea coveredSize numberOfFloors ceilingHeight constructionType
exteriorMaterial roofType foundationType parkingSpaces loadingDocks elevatorCount hvacDetails
fireProtection availabilityStatus propertyStatus availableFor askingAmount possession survey
advance security gracePeriod increment agreementPeriod surveyDate financialIncome
financialExpenses financialMetrics occupancySurvey leaseSurvey utilityDetails units siteDetails
buildingCondition environmentalDetails legalDetails marketDetails attachments createdById
updatedById createdAt updatedAt freshSurveyCount neighboringBusinessIds submittedAt""".split()
EXPECTED_HEADERS = {
    "01_Surveys": ["legacySurveyId"] + SURVEY_FIELDS,
    "02_Survey_Contacts": ["legacySurveyId", "name", "role", "phone", "email", "createdAt"],
    "03_Attachments": ["legacySurveyId", "attachmentType", "originalUrl", "originalFilename",
                       "mimeType", "fileSizeBytes"],
    "04_Exceptions": ["entity", "legacySurveyId", "field", "rawValue", "normalizedValue",
                      "reason", "severity", "notes"],
}
FORBIDDEN_COLUMNS = {
    "id", "migrationStatus", "migrationWarnings", "migrationNotes", "userMatchStatus",
    "businessDeveloperEmail", "matchedUserId", "latitude", "longitude", "Latitude", "Longitude",
    "Business_Developer", "legacyContactSource", "legacyContactIndex", "legacyName", "legacyRole",
    "legacyPhone", "validationStatus", "sourceExists", "targetObjectKey", "checksum",
    "newAttachmentRecordId", "attachmentTag", "legacyAttachmentSource",
}
GRAVITY_SYSTEM = {
    "form_id", "post_id", "is_starred", "is_read", "ip", "source_url", "user_agent", "currency",
    "payment_status", "payment_date", "payment_amount", "payment_method", "transaction_id",
    "is_fulfilled", "created_by", "transaction_type", "Show_Map", "Map_Type",
    "Add_Floor_Wise_Details", "Rent_Sale_Price_Available", "Is_Approved", "status",
    "date_created", "date_updated",
}
PLACEHOLDERS = {"null", "none", "n/a", "undefined", "-"}
ALWAYS_BLANK = [
    "serialNumber", "propertyName", "zipCode", "yearBuilt", "yearRenovated", "buildingClass",
    "ownershipEntity", "propertyManager", "acquisitionDate", "landShape", "topography", "zoning",
    "rentableArea", "usableArea", "ceilingHeight", "constructionType", "exteriorMaterial",
    "roofType", "foundationType", "parkingSpaces", "loadingDocks", "elevatorCount", "hvacDetails",
    "fireProtection", "survey", "financialIncome", "financialExpenses", "financialMetrics",
    "occupancySurvey", "leaseSurvey", "utilityDetails", "siteDetails", "buildingCondition",
    "environmentalDetails", "legalDetails", "attachments", "updatedById", "submittedAt",
]
AREA_COLUMNS = ["Areas_In_Lahore", "Market_Area_114", "Areas_In_Islamabad", "Areas_In_Rawalpindi",
                "Areas_In_Karachi", "Market_Area_117", "Map_Search_Address"]
# approved mapping: destination -> (legacy columns, exception reasons that explain a blank)
APPROVED = {
    "address": (["Address"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "city": (["City"], {"PLACEHOLDER_CITY", "PLACEHOLDER_VALUE_BLANKED"}),
    "area": (AREA_COLUMNS, {"PLACEHOLDER_AREA"}),
    "locationCoordinates": (["Latitude", "Longitude"], {"INVALID_COORDINATES", "MISSING_COORDINATES"}),
    "propertyType": (["Property_Type"], {"AMBIGUOUS_PROPERTY_TYPE"}),
    "ownerName": (["Owner_Name"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "plotSize": (["Plot_Size"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "plotArea": (["Plot_Area"], {"UNPARSEABLE_PLOT_AREA"}),
    "areaUnit": (["Plot_Area"], {"UNPARSEABLE_PLOT_AREA"}),
    "coveredSize": (["Covered_Size"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "grossBuildingArea": (["Total_Area"], {"UNPARSEABLE_TOTAL_AREA"}),
    "numberOfFloors": (["No_of_Floors"], set()),
    "availabilityStatus": (["status"], set()),
    "propertyStatus": (["Available_Not_Rented"], set()),
    "availableFor": (["Available_For"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "askingAmount": (["Amount"], {"UNPARSEABLE_ASKING_AMOUNT"}),
    "advance": (["Advance"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "security": (["Security"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "possession": (["Possession"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "gracePeriod": (["Grace_Period"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "increment": (["Increment"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "agreementPeriod": (["Agreement_Period"], {"PLACEHOLDER_VALUE_BLANKED"}),
    "surveyDate": (["date_created"], {"INVALID_LEGACY_DATE"}),
    "createdAt": (["date_created"], {"INVALID_LEGACY_DATE"}),
    "updatedAt": (["date_updated"], {"INVALID_LEGACY_DATE"}),
    "createdById": (["Business_Developer"], {"BUSINESS_DEVELOPER_MISSING"}),
}
CONTACT_SOURCE = {"name": "Owner_Name", "role": "Contact_Person", "phone": "Contact_Number"}
NUMERIC = ["plotArea", "grossBuildingArea", "askingAmount"]
DIRECT_TEXT = {  # destination: legacy column, compared after trim
    "address": "Address", "ownerName": "Owner_Name", "plotSize": "Plot_Size",
    "coveredSize": "Covered_Size", "possession": "Possession", "advance": "Advance",
    "security": "Security", "gracePeriod": "Grace_Period", "increment": "Increment",
    "agreementPeriod": "Agreement_Period", "availableFor": "Available_For",
    "availabilityStatus": "status",
}
UNIT_KEYS = {"floor", "squareFeet", "dimensions"}


def read_sheet(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return headers, [dict(zip(headers, r)) for r in rows[1:] if any(v is not None for v in r)]


def blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def main():
    errors: list[str] = []
    checks: dict[str, object] = {}

    def check(ok, msg):
        if not ok:
            errors.append(msg)

    with SOURCE.open(encoding="utf-8-sig", newline="") as fh:
        src = list(csv.DictReader(fh))
    by_id = {int(r["id"]): r for r in src}
    bd_emails = {N.s(r["Business_Developer"]).lower() for r in src if "@" in (r["Business_Developer"] or "")}

    wb = load_workbook(XLSX)  # full mode: needed to see formulas/data types
    check(wb.sheetnames == SHEETS, f"Sheets are {wb.sheetnames}, expected {SHEETS}")
    formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f")
    check(formulas == 0, f"{formulas} formula cells found")
    merged = sum(len(ws.merged_cells.ranges) for ws in wb.worksheets)
    check(merged == 0, f"{merged} merged ranges found")
    for ws in wb.worksheets:
        check(ws.freeze_panes == "A2", f"{ws.title}: header row not frozen")
        check(bool(ws.auto_filter.ref), f"{ws.title}: no filter")

    data = {}
    for name in SHEETS:
        headers, rows = read_sheet(wb, name)
        data[name] = rows
        check(headers == EXPECTED_HEADERS[name], f"{name}: headers differ from spec")
        bad = set(headers) & (FORBIDDEN_COLUMNS | GRAVITY_SYSTEM)
        check(not bad, f"{name}: forbidden columns {sorted(bad)}")
        if name != "04_Exceptions":
            ph = [(r.get(headers[0]), h, v) for r in rows for h, v in r.items()
                  if isinstance(v, str) and v.strip().lower() in PLACEHOLDERS]
            check(not ph, f"{name}: literal placeholder values {ph[:5]}")

    # ------------------------------------------------------------ 01_Surveys
    surveys = data["01_Surveys"]
    ids = [r["legacySurveyId"] for r in surveys]
    check(len(surveys) == 7398, f"01_Surveys rows={len(surveys)}")
    check(len(set(ids)) == len(ids), "legacySurveyId not unique")
    check(set(ids) == set(by_id), "legacySurveyId set differs from source")
    survey_ids = set(ids)
    json_counts = Counter()
    created_by_mismatch = 0
    for r in surveys:
        lid = r["legacySurveyId"]
        s = by_id[lid]
        for f in ALWAYS_BLANK:
            check(blank(r[f]), f"{lid}: {f} should be blank")
        for f in NUMERIC:
            check(blank(r[f]) or isinstance(r[f], (int, float)), f"{lid}: {f} not numeric")
        check(blank(r["numberOfFloors"]) or isinstance(r["numberOfFloors"], int), f"{lid}: numberOfFloors not int")
        check(r["freshSurveyCount"] == 0, f"{lid}: freshSurveyCount != 0")
        check(r["neighboringBusinessIds"] == "[]", f"{lid}: neighboringBusinessIds != []")
        # dates
        ca, sd = r["createdAt"], r["surveyDate"]
        check(blank(ca) or isinstance(ca, datetime), f"{lid}: createdAt invalid")
        check(blank(sd) or isinstance(sd, (date, datetime)), f"{lid}: surveyDate invalid")
        check(blank(r["updatedAt"]) or isinstance(r["updatedAt"], datetime), f"{lid}: updatedAt invalid")
        if isinstance(ca, datetime):
            check(sd is not None and sd.date() == ca.date(), f"{lid}: surveyDate != createdAt date")
            check(ca.strftime("%Y-%m-%d %H:%M:%S") == s["date_created"], f"{lid}: createdAt != date_created")
        # json
        if not blank(r["locationCoordinates"]):
            c = json.loads(r["locationCoordinates"])
            check(set(c) == {"latitude", "longitude"} and all(isinstance(v, (int, float)) for v in c.values()),
                  f"{lid}: bad locationCoordinates")
            json_counts["locationCoordinates"] += 1
        if not blank(r["units"]):
            u = json.loads(r["units"])
            check(isinstance(u, list) and u and all(isinstance(x, dict) and set(x) <= UNIT_KEYS and x for x in u),
                  f"{lid}: bad units")
            json_counts["units"] += 1
        if not blank(r["marketDetails"]):
            m = json.loads(r["marketDetails"])
            check(set(m) == {"nearbyCompetitors"} and all(isinstance(x, str) and x for x in m["nearbyCompetitors"]),
                  f"{lid}: bad marketDetails")
            json_counts["marketDetails"] += 1
        json.loads(r["neighboringBusinessIds"])
        # BD email only in createdById, and only this record's own value
        for f, v in r.items():
            if f != "createdById" and isinstance(v, str) and (v.lower() in bd_emails or "@shapers.com.pk" in v.lower()):
                errors.append(f"{lid}: Business_Developer email in {f}")
        bd = N.s(s["Business_Developer"])
        expected_bd = bd if "@" in bd else None
        if (r["createdById"] or None) != expected_bd:
            created_by_mismatch += 1
            errors.append(f"{lid}: createdById {r['createdById']!r} != own Business_Developer {bd!r}")
        # direct mappings match the source
        for dest, col in DIRECT_TEXT.items():
            if not blank(r[dest]):
                check(r[dest] == N.s(s[col]), f"{lid}: {dest} differs from {col}")
        amt, _ = N.parse_asking_amount(s["Amount"])
        check((amt is None and blank(r["askingAmount"])) or (amt is not None and float(amt) == r["askingAmount"]),
              f"{lid}: askingAmount differs from Amount")
    checks["jsonCellsParsed"] = dict(json_counts)
    filled = [r["createdById"] for r in surveys if not blank(r["createdById"])]
    checks["createdById"] = {
        "source": "Business_Developer", "matchingKey": "legacySurveyId",
        "withEmail": len(filled), "blank": len(surveys) - len(filled),
        "distinctEmails": len(set(filled)),
        "excludedPlaceholders": dict(Counter(N.s(by_id[r["legacySurveyId"]]["Business_Developer"]) or "(blank)"
                                             for r in surveys if blank(r["createdById"]))),
        "crossRecordMismatches": created_by_mismatch,
    }

    # ------------------------------------------------------------ exceptions index
    exc = data["04_Exceptions"]
    for e in exc:
        check(e["entity"] in {"Dataset", "Survey", "SurveyContact", "Attachment"}, f"bad entity {e['entity']}")
        if e["entity"] != "Dataset":
            check(e["legacySurveyId"] in survey_ids, f"exception for unknown survey {e['legacySurveyId']}")
    exc_by = defaultdict(set)
    for e in exc:
        exc_by[(e["reason"], e["field"])].add(e["legacySurveyId"])
    reasons = Counter(e["reason"] for e in exc)
    check(reasons["LEGACY_STATUS_CONSTANT_ACTIVE"] == 1, "LEGACY_STATUS_CONSTANT_ACTIVE row missing")

    # ------------------------------------------------------------ approved-mapping gaps
    # blank destination + same-record source value => must be explained by an exception
    reasons_by_id = defaultdict(set)
    for e in exc:
        reasons_by_id[e["legacySurveyId"]].add(e["reason"])
    gaps = {}
    for dest, (cols, ok_reasons) in APPROVED.items():
        found = [r["legacySurveyId"] for r in surveys
                 if blank(r[dest]) and any(N.s(by_id[r["legacySurveyId"]][c]) for c in cols)]
        unexplained = [i for i in found if not (reasons_by_id[i] & ok_reasons)]
        if found:
            gaps[dest] = {"blankWithSource": len(found), "unexplained": len(unexplained)}
        check(not unexplained, f"{dest}: {len(unexplained)} blank cells have an unexplained same-record source value, e.g. {unexplained[:5]}")
    checks["approvedMappingGaps"] = gaps

    # ------------------------------------------------------------ preservation
    pres = {}
    for field in ("No_Rent_Price_Text", "Deal", "Note", "Area_Details_Box"):
        populated = {int(r["id"]) for r in src if N.s(r[field])}
        got = {e["legacySurveyId"] for e in exc if e["field"] == field and e["reason"] == "PENDING_BUSINESS_DECISION"}
        raw_ok = all(e["rawValue"] == N.s(by_id[e["legacySurveyId"]][field])
                     for e in exc if e["field"] == field and e["reason"] == "PENDING_BUSINESS_DECISION")
        check(populated == got and raw_ok, f"{field}: pending values not all preserved")
        pres[field] = {"populated": len(populated), "inExceptions": len(got)}
    sv = {r["legacySurveyId"]: r for r in surveys}
    for col, dest, reason in (("Total_Area", "grossBuildingArea", "UNPARSEABLE_TOTAL_AREA"),
                              ("Plot_Area", "plotArea", "UNPARSEABLE_PLOT_AREA")):
        populated = {int(r["id"]) for r in src if N.s(r[col])}
        written = {i for i in populated if not blank(sv[i][dest])}
        flagged = exc_by[(reason, dest)]
        check(written | flagged == populated and not (written & flagged), f"{col}: values neither written nor flagged")
        pres[col] = {"populated": len(populated), "written": len(written), "inExceptions": len(flagged)}
    bad_dates = {int(r["id"]) for r in src if N.parse_date(r["date_created"]) is None}
    check(bad_dates == exc_by[("INVALID_LEGACY_DATE", "surveyDate/createdAt")], "invalid dates not all flagged")
    check(all(blank(sv[i]["createdAt"]) and blank(sv[i]["surveyDate"]) for i in bad_dates), "invalid dates were filled")
    pres["date_created"] = {"valid": len(src) - len(bad_dates), "inExceptions": len(bad_dates)}
    unmapped_pt = {i for i, r in sv.items() if blank(r["propertyType"])}
    flagged_pt = exc_by[("AMBIGUOUS_PROPERTY_TYPE", "propertyType")] | exc_by[("MISSING_PROPERTY_TYPE", "propertyType")]
    check(unmapped_pt == flagged_pt, "blank propertyType not all flagged")
    pres["propertyType"] = {"mapped": len(sv) - len(unmapped_pt), "inExceptions": len(flagged_pt)}
    conflicts = exc_by[("CONFLICT_WITH_POSSESSION", "propertyStatus")]
    check(all(not blank(sv[i]["propertyStatus"]) for i in conflicts), "conflict rows lost propertyStatus")
    pres["propertyStatusConflicts"] = len(conflicts)

    # floors: every used legacy slot is a unit
    lost_slots = 0
    top_slot = 0
    for r in src:
        used = [i for i in range(1, 12) if any(N.s(r.get(f"Floor_{i}_{k}")) for k in ("Type", "Area", "Size"))]
        if used:
            top_slot = max(top_slot, max(used))
            units = json.loads(sv[int(r["id"])]["units"] or "[]")
            lost_slots += max(0, len(used) - len(units))
    check(lost_slots == 0, f"{lost_slots} floor slots missing from units")
    pres["floorSlotsLost"] = lost_slots
    pres["highestFloorSlot"] = top_slot

    # brands: every legacy brand is in nearbyCompetitors
    lost_brands = 0
    for r in src:
        comp = {x.lower() for x in json.loads(sv[int(r["id"])]["marketDetails"] or '{"nearbyCompetitors":[]}')["nearbyCompetitors"]}
        for part in N.s(r["Select_Brands"]).split(" | "):
            for b in N.parse_json_list(part):
                if isinstance(b, str) and N.s(b) and N.s(b).lower() not in comp:
                    lost_brands += 1
        for col in ("Other_Brands", "Earlier_Brands"):
            text = N.s(r[col])
            if text:
                core = re.sub(r"[,\s]*\betc\b\.?\s*$", "", text, flags=re.I).strip(" ,;.")
                pieces = [p.strip().lower() for p in re.split(r"[,;\n]", core) if p.strip()]
                if text.lower() not in comp and not all(p in comp for p in pieces):
                    lost_brands += 1
    check(lost_brands == 0, f"{lost_brands} legacy brand values missing from marketDetails")
    pres["brandValuesLost"] = lost_brands

    # ------------------------------------------------------------ contacts
    contacts = data["02_Survey_Contacts"]
    classes = {int(r["id"]): N.classify_contact(r) for r in src}
    hold = {i for i, c in classes.items() if c["migrationStatus"] == "HOLD"}
    cids = [c["legacySurveyId"] for c in contacts]
    check(len(contacts) == len(src), f"02_Survey_Contacts rows={len(contacts)}, expected {len(src)}")
    check(len(set(cids)) == len(cids) and set(cids) == survey_ids,
          "contacts are not exactly one row per 01_Surveys legacySurveyId")
    blanked = {(e["legacySurveyId"], e["field"]) for e in exc
               if e["entity"] == "SurveyContact" and e["reason"] == "PLACEHOLDER_VALUE_BLANKED"}
    mismatches = Counter()
    for c in contacts:
        lid = c["legacySurveyId"]
        s = by_id[lid]
        for dest, col in CONTACT_SOURCE.items():
            raw = N.s(s[col])
            got = c[dest] if not blank(c[dest]) else ""
            if got != raw and not (got == "" and (lid, dest) in blanked):
                mismatches[dest] += 1
                errors.append(f"contact {lid}: {dest} {got!r} != own {col} {raw!r}")
        check(blank(c["email"]), f"contact {lid}: email not blank")
        check(c["createdAt"] == sv[lid]["createdAt"], f"contact {lid}: createdAt differs from survey")
    hold_exc = {e["legacySurveyId"] for e in exc if e["entity"] == "SurveyContact" and e["field"] == "contact"}
    check(hold_exc == hold, "previously HOLD contacts not all kept as exceptions")
    check(hold <= set(cids), "previously HOLD contact missing from 02_Survey_Contacts")
    cmap = {c["legacySurveyId"]: c for c in contacts}
    has = lambda f: sum(1 for c in contacts if not blank(c[f]))  # noqa: E731
    phone_warn = {e["legacySurveyId"] for e in exc if e["reason"] == "CONTACT_PHONE_INCOMPLETE_OR_INVALID"}
    multi = {e["legacySurveyId"] for e in exc if e["reason"] == "CONTACT_MULTIPLE_PHONES"}
    checks["contacts"] = {
        "rows": len(contacts),
        "withName": has("name"), "withoutName": len(contacts) - has("name"),
        "withRole": has("role"), "withoutRole": len(contacts) - has("role"),
        "withPhone": has("phone"), "withoutPhone": len(contacts) - has("phone"),
        "incompleteOrInvalidPhone": len(phone_warn),
        "multiplePhones": len(multi),
        "multiplePhonesKeptInFull": sum(1 for i in multi if cmap[i]["phone"] == N.s(by_id[i]["Contact_Number"])),
        "invalidName": sum(1 for e in exc if e["reason"] == "CONTACT_NAME_INVALID"),
        "previouslyHold": len(hold), "previouslyHoldPresent": len(hold & set(cids)),
        "previousHoldReasons": dict(Counter(classes[i]["validationStatus"] for i in hold)),
        "placeholdersBlanked": len(blanked),
        "sourceMismatches": dict(mismatches), "crossRecordMismatches": sum(mismatches.values()),
    }

    # ------------------------------------------------------------ attachments
    atts = data["03_Attachments"]
    head = N.load_head_scan(HEAD)
    all_refs = []
    for r in src:
        all_refs.extend(N.extract_attachments(r, head))
    ready = {(a["legacySurveyId"], a["originalUrl"]) for a in all_refs if a["migrationStatus"] == "READY"}
    got = {(a["legacySurveyId"], a["originalUrl"]) for a in atts}
    check(got <= ready, "non-READY attachment in clean sheet")
    check(len(got) == len(atts), "duplicate attachment rows")
    for a in atts:
        check(a["attachmentType"] in ("IMAGE", "VIDEO"), f"attachment type {a['attachmentType']}")
        check(not a["originalUrl"].rstrip("/").endswith("/1"), "placeholder /1 attachment in clean sheet")
        check(a["legacySurveyId"] in survey_ids, "attachment references unknown survey")
        check(blank(a["fileSizeBytes"]) or isinstance(a["fileSizeBytes"], int), "fileSizeBytes not int")
    att_exc = [e for e in exc if e["entity"] == "Attachment"]
    check(len(atts) + len(att_exc) == len(all_refs), f"attachments {len(atts)} + exceptions {len(att_exc)} != {len(all_refs)}")
    checks["attachments"] = {"clean": len(atts), "types": dict(Counter(a["attachmentType"] for a in atts)),
                             "inExceptions": len(att_exc), "legacyReferences": len(all_refs)}

    report = {
        "ok": not errors,
        "errorCount": len(errors),
        "errors": errors[:200],
        "rows": {name: len(rows) for name, rows in data.items()},
        "surveyColumns": len(EXPECTED_HEADERS["01_Surveys"]),
        "checks": checks,
        "preservation": pres,
        "exceptionsByReason": dict(reasons.most_common()),
        "workbook": XLSX.name,
    }
    REPORT.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    print(json.dumps(report, indent=2, default=str))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    if len(sys.argv) > 1:  # optional workbook path
        XLSX = Path(sys.argv[1]).resolve()
    main()
