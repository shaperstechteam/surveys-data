"""Build 00_Final_Schema_Mapping — two-sided final field ← legacy source index."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def coverage(rows, col):
    return sum(1 for r in rows if (r.get(col) or "").strip())

FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FONT_H = Font(name="Arial", size=14, bold=True)
FONT_W = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
NEW_FILL = PatternFill("solid", fgColor="C6EFCE")
EXISTING_FILL = PatternFill("solid", fgColor="DDEBF7")
NONE_FILL = PatternFill("solid", fgColor="F2F2F2")
RAW_FILL = PatternFill("solid", fgColor="E2D5F1")
PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="833C0C")
WRAP = Alignment(wrap_text=True, vertical="top")

HEADERS = [
    "Entity",
    "FinalField",
    "FinalType",
    "ExistingOrNew",
    "LegacySourceColumn",
    "LegacySourceMeaning",
    "MappingRule",
    "PopulatedLegacyRows",
    "ExampleLegacyValue",
    "FinalExample",
    "Decision",
    "Notes",
]

APPROVED_NEW = {
    "area",
    "plotSize",
    "coveredSize",
    "availableFor",
    "askingAmount",
    "possession",
    "survey",
    "advance",
    "security",
    "gracePeriod",
    "increment",
    "agreementPeriod",
}

NO_SOURCE = "NO LEGACY SOURCE"
KEEP_NULL = "KEEP CURRENT SCHEMA / NULL FOR LEGACY IMPORT"


def _ex(rows, col, n=1):
    for r in rows:
        v = (r.get(col) or "").strip()
        if v:
            return v[:220]
    return ""


def _row(**kwargs):
    return {h: kwargs.get(h, "") for h in HEADERS}


def survey_rows(rows: list[dict], surveys: list[dict]) -> list[dict]:
    first = next((s for s in surveys if s.get("locationCoordinates")), surveys[0])
    area_ex = next((s for s in surveys if s.get("area")), first)
    amt_ex = next((s for s in surveys if s.get("askingAmount") not in ("", None)), first)

    def pop(*cols):
        n = 0
        for r in rows:
            if any((r.get(c) or "").strip() for c in cols):
                n += 1
        return n

    none = lambda field, typ, notes="": _row(
        Entity="Survey",
        FinalField=field,
        FinalType=typ,
        ExistingOrNew="EXISTING",
        LegacySourceColumn=NO_SOURCE,
        LegacySourceMeaning="",
        MappingRule="Leave null / application default",
        PopulatedLegacyRows=0,
        ExampleLegacyValue="",
        FinalExample="",
        Decision=KEEP_NULL,
        Notes=notes,
    )

    out = [
        _row(
            Entity="Survey",
            FinalField="id",
            FinalType="Int (generated)",
            ExistingOrNew="EXISTING",
            LegacySourceColumn=NO_SOURCE,
            LegacySourceMeaning="Legacy id is NOT this field",
            MappingRule="Database autoincrement after import",
            PopulatedLegacyRows=0,
            Decision=KEEP_NULL,
            Notes="See MIGRATION METADATA: id → legacySurveyId",
        ),
        _row(
            Entity="Survey",
            FinalField="serialNumber",
            FinalType="String @unique",
            ExistingOrNew="EXISTING",
            LegacySourceColumn=NO_SOURCE,
            MappingRule="Generate via CRE SerialNumberService at import",
            PopulatedLegacyRows=0,
            FinalExample="(blank in workbook)",
            Decision=KEEP_NULL,
            Notes="Do not reuse legacy id as serialNumber",
        ),
        none("propertyName", "String?"),
        _row(
            Entity="Survey",
            FinalField="address",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Address",
            LegacySourceMeaning="Street address",
            MappingRule="Direct trim",
            PopulatedLegacyRows=coverage(rows, "Address"),
            ExampleLegacyValue=_ex(rows, "Address"),
            FinalExample=first.get("address") or "",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="locationCoordinates",
            FinalType="Json? (CRE today is still Decimal(9,6))",
            ExistingOrNew="EXISTING — TYPE CHANGE REQUIRED ON CRE",
            LegacySourceColumn="Latitude + Longitude",
            LegacySourceMeaning="WGS84 pair",
            MappingRule='One JSON object {"latitude":…,"longitude":…}. No permanent lat/lng Survey columns.',
            PopulatedLegacyRows=pop("Latitude", "Longitude"),
            ExampleLegacyValue=f"{_ex(rows,'Latitude')} , {_ex(rows,'Longitude')}",
            FinalExample=first.get("locationCoordinates") or "",
            Decision="APPROVED",
            Notes="Workbook preserves both values. CRE must change Decimal → Json before import.",
        ),
        _row(
            Entity="Survey",
            FinalField="city",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="City",
            LegacySourceMeaning="City name",
            MappingRule="Direct; spelling map; Select/. → null",
            PopulatedLegacyRows=coverage(rows, "City"),
            ExampleLegacyValue=_ex(rows, "City"),
            FinalExample=area_ex.get("city") or "",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="area",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Areas_In_Lahore / Market_Area_114 / Areas_In_Islamabad / Areas_In_Rawalpindi / Areas_In_Karachi / Market_Area_117 / Map_Search_Address fallback",
            LegacySourceMeaning="City-specific locality then market then maps address",
            MappingRule="Exactly one string. City=Lahore→Areas_In_Lahore; Islamabad→Areas_In_Islamabad; Rawalpindi→Areas_In_Rawalpindi; Karachi→Areas_In_Karachi; else Areas_In_Karachi then markets; Map_Search_Address last.",
            PopulatedLegacyRows=sum(1 for s in surveys if s.get("area")),
            ExampleLegacyValue=_ex(rows, "Areas_In_Islamabad") or _ex(rows, "Areas_In_Lahore"),
            FinalExample=area_ex.get("area") or "",
            Decision="APPROVED",
            Notes="Not an area JSON. Not a second city field.",
        ),
        none("zipCode", "String?"),
        none("yearBuilt", "Int?"),
        none("yearRenovated", "Int?"),
        none("buildingClass", "BuildingClass?"),
        _row(
            Entity="Survey",
            FinalField="propertyType",
            FinalType="PropertyType?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Property_Type",
            LegacySourceMeaning="building | plot | office | house",
            MappingRule="plot→LAND_PLOT; office→OFFICE; house→RESIDENTIAL; building/blank→null + REVIEW",
            PopulatedLegacyRows=coverage(rows, "Property_Type"),
            ExampleLegacyValue="building",
            FinalExample="",
            Decision="APPROVED",
            Notes="Do not force building → RETAIL",
        ),
        _row(
            Entity="Survey",
            FinalField="ownerName",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Owner_Name",
            LegacySourceMeaning="Owner / company",
            MappingRule="Direct trim; also SurveyContact.name when valid",
            PopulatedLegacyRows=coverage(rows, "Owner_Name"),
            ExampleLegacyValue=_ex(rows, "Owner_Name"),
            FinalExample=first.get("ownerName") or "",
            Decision="APPROVED",
        ),
        none("ownershipEntity", "String?"),
        none("propertyManager", "String?"),
        none("acquisitionDate", "Date?"),
        _row(
            Entity="Survey",
            FinalField="plotArea",
            FinalType="Decimal(12,2)?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Plot_Area",
            LegacySourceMeaning="Native-unit plot area",
            MappingRule="Parse only when unit is deterministic; else null + flag",
            PopulatedLegacyRows=coverage(rows, "Plot_Area"),
            ExampleLegacyValue=_ex(rows, "Plot_Area"),
            FinalExample=str(first.get("plotArea") or ""),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="areaUnit",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Plot_Area",
            LegacySourceMeaning="Unit inside Plot_Area text",
            MappingRule="Kanal / Marla / Acre / Sq Ft / Sq Yd when parseable",
            PopulatedLegacyRows=coverage(rows, "Plot_Area"),
            ExampleLegacyValue=_ex(rows, "Plot_Area"),
            FinalExample=first.get("areaUnit") or "",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="plotSize",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Plot_Size",
            LegacySourceMeaning="Plot dimension string",
            MappingRule="Preserve original dimension text; do not convert to a number",
            PopulatedLegacyRows=coverage(rows, "Plot_Size"),
            ExampleLegacyValue=_ex(rows, "Plot_Size"),
            FinalExample=_ex(rows, "Plot_Size"),
            Decision="APPROVED",
        ),
        none("landShape", "String?"),
        none("topography", "String?"),
        none("zoning", "String?"),
        _row(
            Entity="Survey",
            FinalField="grossBuildingArea",
            FinalType="Decimal(12,2)?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Total_Area",
            MappingRule="Parse clean sq ft only; approx flagged",
            PopulatedLegacyRows=coverage(rows, "Total_Area"),
            ExampleLegacyValue=_ex(rows, "Total_Area"),
            FinalExample=str(amt_ex.get("grossBuildingArea") or ""),
            Decision="APPROVED",
        ),
        none("rentableArea", "Decimal?"),
        none("usableArea", "Decimal?"),
        _row(
            Entity="Survey",
            FinalField="coveredSize",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Covered_Size",
            LegacySourceMeaning="Covered dimension string",
            MappingRule="Preserve original dimension text",
            PopulatedLegacyRows=coverage(rows, "Covered_Size"),
            ExampleLegacyValue=_ex(rows, "Covered_Size"),
            FinalExample=_ex(rows, "Covered_Size"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="numberOfFloors",
            FinalType="Int?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="No_of_Floors",
            MappingRule="Integral only",
            PopulatedLegacyRows=coverage(rows, "No_of_Floors"),
            ExampleLegacyValue=_ex(rows, "No_of_Floors"),
            FinalExample=str(first.get("numberOfFloors") or ""),
            Decision="APPROVED",
        ),
        none("ceilingHeight", "Decimal?"),
        none("constructionType", "String?"),
        none("exteriorMaterial", "String?"),
        none("roofType", "String?"),
        none("foundationType", "String?"),
        none("parkingSpaces", "Int?"),
        none("loadingDocks", "Int?"),
        none("elevatorCount", "Int?"),
        none("hvacDetails", "String?"),
        none("fireProtection", "String?"),
        _row(
            Entity="Survey",
            FinalField="availabilityStatus",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="status",
            LegacySourceMeaning="Gravity Forms entry status",
            MappingRule="Preserve legacy status string. NOT propertyStatus.",
            PopulatedLegacyRows=coverage(rows, "status"),
            ExampleLegacyValue="active",
            FinalExample="active",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="propertyStatus",
            FinalType="PropertyStatus?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Available_Not_Rented",
            LegacySourceMeaning="Yes / No checkbox",
            MappingRule="Yes→VACANT; No→RENTED; other→null. Conflict with Possession → REVIEW",
            PopulatedLegacyRows=coverage(rows, "Available_Not_Rented"),
            ExampleLegacyValue="Yes",
            FinalExample="VACANT",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="availableFor",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Available_For",
            LegacySourceMeaning="Rent / Sale",
            MappingRule="Preserve/normalize casing only if safe; Select/. → null",
            PopulatedLegacyRows=coverage(rows, "Available_For"),
            ExampleLegacyValue=_ex(rows, "Available_For"),
            FinalExample="Rent",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="askingAmount",
            FinalType="Decimal(14,2)?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Amount",
            LegacySourceMeaning="Asking rent or sale PKR",
            MappingRule="Parse numeric (commas/currency stripped). Unparseable → null + REVIEW",
            PopulatedLegacyRows=coverage(rows, "Amount"),
            ExampleLegacyValue=_ex(rows, "Amount"),
            FinalExample=str(amt_ex.get("askingAmount") or ""),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="possession",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Possession",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Possession"),
            ExampleLegacyValue=_ex(rows, "Possession"),
            FinalExample=_ex(rows, "Possession"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="survey",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="NO CONFIRMED SOURCE",
            MappingRule="Do not invent",
            PopulatedLegacyRows=0,
            ExampleLegacyValue="",
            FinalExample="",
            Decision="APPROVED FIELD / LEGACY VALUE NULL",
            Notes="Approved Prisma field; 96-column source has no corresponding column.",
        ),
        _row(
            Entity="Survey",
            FinalField="advance",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Advance",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Advance"),
            ExampleLegacyValue=_ex(rows, "Advance"),
            FinalExample=_ex(rows, "Advance"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="security",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Security",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Security"),
            ExampleLegacyValue=_ex(rows, "Security"),
            FinalExample=_ex(rows, "Security"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="gracePeriod",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Grace_Period",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Grace_Period"),
            ExampleLegacyValue=_ex(rows, "Grace_Period"),
            FinalExample=_ex(rows, "Grace_Period"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="increment",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Increment",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Increment"),
            ExampleLegacyValue=_ex(rows, "Increment"),
            FinalExample=_ex(rows, "Increment"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="agreementPeriod",
            FinalType="String?",
            ExistingOrNew="NEW",
            LegacySourceColumn="Agreement_Period",
            MappingRule="Preserve original string",
            PopulatedLegacyRows=coverage(rows, "Agreement_Period"),
            ExampleLegacyValue=_ex(rows, "Agreement_Period"),
            FinalExample=_ex(rows, "Agreement_Period"),
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="surveyDate",
            FinalType="DateTime (CRE currently NOT NULL)",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="date_created",
            MappingRule="Valid dates only; 0000-00-00 → null",
            PopulatedLegacyRows=sum(1 for r in rows if (r.get("date_created") or "").strip() and not str(r.get("date_created")).startswith("0000-00-00")),
            ExampleLegacyValue=_ex(rows, "date_created"),
            FinalExample=first.get("surveyDate") or "",
            Decision="APPROVED",
            Notes="1,491 zero dates. CRE surveyDate must be nullable before those rows can import.",
        ),
        none("financialIncome", "Json?", "Asking Amount is askingAmount, not income"),
        none("financialExpenses", "Json?"),
        none("financialMetrics", "Json?", "Server-computed; do not populate from legacy"),
        none("occupancySurvey", "Json?", "Server-computed from units/leaseSurvey"),
        none("leaseSurvey", "Json?", "Offer terms are not an existing lease"),
        none("utilityDetails", "Json?"),
        _row(
            Entity="Survey",
            FinalField="units",
            FinalType="Json?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Select_Floors + Floor_1_Type/Area/Size … Floor_8_Type/Area/Size + Floor_9_Type/Area + Floor_10_Type/Area + Floor_11_Type/Area",
            LegacySourceMeaning="Floor labels, areas, dimensions through floor 11",
            MappingRule='Array of {floor, squareFeet?, dimensions?}. Dimensions are IN the Property Viewer units JSON. CRE SurveyUnitDto does not yet allow dimensions — extend DTO later.',
            PopulatedLegacyRows=sum(1 for s in surveys if s.get("units")),
            ExampleLegacyValue=_ex(rows, "Floor_1_Type"),
            FinalExample='{"floor":"Ground Floor","squareFeet":2700,"dimensions":"032ft x 085ft"}',
            Decision="APPROVED",
            Notes="Area_Details_Box is NOT copied into units (see PENDING). Floor 1–11 not truncated.",
        ),
        none("siteDetails", "Json?"),
        none("buildingCondition", "Json?"),
        none("environmentalDetails", "Json?"),
        none("legalDetails", "Json?"),
        _row(
            Entity="Survey",
            FinalField="marketDetails",
            FinalType="Json?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Earlier_Brands + Select_Brands + Other_Brands",
            MappingRule="nearbyCompetitors string[]. neighboringBusinessIds stays [] until CRE Business.id lookup",
            PopulatedLegacyRows=sum(1 for s in surveys if s.get("marketDetails")),
            ExampleLegacyValue=_ex(rows, "Select_Brands")[:120],
            FinalExample=(first.get("marketDetails") or "")[:180],
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="attachments",
            FinalType="Json?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Pictures + Videos + Backup_Documents",
            MappingRule="Blank on 02 until S3 upload. Manifest on 04_Attachments.",
            PopulatedLegacyRows=pop("Pictures", "Videos", "Backup_Documents"),
            ExampleLegacyValue="(JSON URL arrays — see 04)",
            FinalExample="",
            Decision="APPROVED",
            Notes="CRE AttachmentRecord is IMAGE|VIDEO only; OTHER backup files HOLD.",
        ),
        _row(
            Entity="Survey",
            FinalField="createdById",
            FinalType="Int?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Business_Developer",
            LegacySourceMeaning="Internal Shapers BD email — NOT a property-contact email",
            MappingRule="trim+lowercase email → CRE User.email → User.id. Workbook leaves createdById/matchedUserId null.",
            PopulatedLegacyRows=coverage(rows, "Business_Developer"),
            ExampleLegacyValue=_ex(rows, "Business_Developer"),
            FinalExample="null until CRE User lookup",
            Decision="APPROVED",
            Notes="NEVER SurveyContact.email. Gravity Forms created_by is ignored.",
        ),
        none("updatedById", "Int?"),
        _row(
            Entity="Survey",
            FinalField="createdAt",
            FinalType="DateTime",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="date_created",
            MappingRule="Same valid date as surveyDate; else importer/DB default",
            PopulatedLegacyRows=sum(1 for r in rows if (r.get("date_created") or "").strip() and not str(r.get("date_created")).startswith("0000-00-00")),
            ExampleLegacyValue=_ex(rows, "date_created"),
            FinalExample=first.get("createdAt") or "",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="updatedAt",
            FinalType="DateTime",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="date_updated",
            MappingRule="Valid historical value only. Source is empty/0000 for this extract → blank.",
            PopulatedLegacyRows=sum(1 for r in rows if (r.get("date_updated") or "").strip() and not str(r.get("date_updated")).startswith("0000")),
            ExampleLegacyValue=_ex(rows, "date_updated") or "(empty in source)",
            FinalExample="",
            Decision="APPROVED",
        ),
        _row(
            Entity="Survey",
            FinalField="freshSurveyCount",
            FinalType="Int",
            ExistingOrNew="EXISTING",
            LegacySourceColumn=NO_SOURCE,
            MappingRule="System default 0",
            PopulatedLegacyRows=0,
            FinalExample="0",
            Decision=KEEP_NULL,
        ),
        _row(
            Entity="Survey",
            FinalField="neighboringBusinessIds",
            FinalType="Int[]",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Select_Brands / Earlier_Brands / Other_Brands (candidates only)",
            MappingRule="Do not write fake ids. [] until CRE Business lookup. Names go to marketDetails.",
            PopulatedLegacyRows=0,
            FinalExample="[]",
            Decision="APPROVED",
        ),
        none("submittedAt", "DateTime?"),
    ]
    return out


def contact_rows(rows: list[dict]) -> list[dict]:
    return [
        _row(
            Entity="SurveyContact",
            FinalField="id",
            FinalType="Int (generated)",
            ExistingOrNew="EXISTING",
            LegacySourceColumn=NO_SOURCE,
            MappingRule="Database generated after import",
            Decision=KEEP_NULL,
        ),
        _row(
            Entity="SurveyContact",
            FinalField="surveyId",
            FinalType="Int",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="legacySurveyId → new Survey.id",
            MappingRule="Importer resolves after Survey insert",
            Decision="APPROVED",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="name",
            FinalType="String (required)",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Owner_Name",
            MappingRule="Only when name is valid. Never fabricate Unknown/Owner/N/A.",
            PopulatedLegacyRows=coverage(rows, "Owner_Name"),
            ExampleLegacyValue=_ex(rows, "Owner_Name"),
            Decision="APPROVED",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="role",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Contact_Person",
            MappingRule="Pure role words only (Owner, Caretaker, …)",
            PopulatedLegacyRows=coverage(rows, "Contact_Person"),
            ExampleLegacyValue=_ex(rows, "Contact_Person"),
            Decision="APPROVED",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="phone",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="Contact_Number",
            MappingRule="One safe PK number → +92…; multi-phone HOLD; do not split",
            PopulatedLegacyRows=coverage(rows, "Contact_Number"),
            ExampleLegacyValue=_ex(rows, "Contact_Number"),
            Decision="APPROVED",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="email",
            FinalType="String?",
            ExistingOrNew="EXISTING",
            LegacySourceColumn=NO_SOURCE,
            MappingRule="No property-contact email in legacy source. Always null.",
            PopulatedLegacyRows=0,
            FinalExample="",
            Decision=KEEP_NULL,
            Notes="Business_Developer → Survey.createdById. NOT SurveyContact.email.",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="createdAt",
            FinalType="DateTime",
            ExistingOrNew="EXISTING",
            LegacySourceColumn="date_created",
            MappingRule="Valid date_created or importer/DB default",
            Decision="APPROVED",
        ),
        _row(
            Entity="SurveyContact",
            FinalField="(explicit non-mapping)",
            FinalType="",
            ExistingOrNew="N/A",
            LegacySourceColumn="Business_Developer",
            MappingRule="→ Survey.createdById via User.email. NOT SurveyContact.email.",
            PopulatedLegacyRows=coverage(rows, "Business_Developer"),
            ExampleLegacyValue=_ex(rows, "Business_Developer"),
            Decision="APPROVED",
            Notes="Do not copy BD into contact email.",
        ),
    ]


def raw_rows(rows: list[dict]) -> list[dict]:
    fields = [
        ("id", "Gravity Forms entry ID", "MIGRATION METADATA ONLY → legacySurveyId. Never Survey.id."),
        ("form_id", "GF form id", "RAW / SYSTEM — DO NOT IMPORT"),
        ("post_id", "WP post id (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("is_starred", "GF starred", "RAW / SYSTEM — DO NOT IMPORT"),
        ("is_read", "GF read", "RAW / SYSTEM — DO NOT IMPORT"),
        ("ip", "Submitter IP", "RAW / SYSTEM — DO NOT IMPORT"),
        ("source_url", "Form URL", "RAW / SYSTEM — DO NOT IMPORT"),
        ("user_agent", "Browser UA", "RAW / SYSTEM — DO NOT IMPORT"),
        ("currency", "GF payment", "RAW / SYSTEM — DO NOT IMPORT"),
        ("payment_status", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("payment_date", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("payment_amount", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("payment_method", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("transaction_id", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("is_fulfilled", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("created_by", "WP user id of the clerk", "RAW / SYSTEM — DO NOT IMPORT. Not createdById."),
        ("transaction_type", "GF payment (empty)", "RAW / SYSTEM — DO NOT IMPORT"),
        ("Show_Map", "Map widget mode", "RAW / SYSTEM — DO NOT IMPORT"),
        ("Map_Type", "Roadmap / Satellite", "RAW / SYSTEM — DO NOT IMPORT"),
        ("Add_Floor_Wise_Details", "UI switch Yes/No/Details Box", "RAW / SYSTEM — DO NOT IMPORT (floors themselves → units)"),
        ("Rent_Sale_Price_Available", "Equals Amount filled", "RAW / SYSTEM — DO NOT IMPORT"),
        ("Is_Approved", "GF approval code", "RAW / SYSTEM — DO NOT IMPORT"),
    ]
    out = []
    for col, meaning, decision in fields:
        entity = "MIGRATION METADATA ONLY" if col == "id" else "RAW / SYSTEM — DO NOT IMPORT"
        dest = "legacySurveyId" if col == "id" else ""
        out.append(
            _row(
                Entity=entity,
                FinalField=dest,
                FinalType="",
                ExistingOrNew="RAW-ONLY" if col != "id" else "MIGRATION METADATA",
                LegacySourceColumn=col,
                LegacySourceMeaning=meaning,
                MappingRule="Preserve on 01_Legacy_Raw_Data only" if col != "id" else "Workbook helper + importer manifest",
                PopulatedLegacyRows=coverage(rows, col),
                ExampleLegacyValue=_ex(rows, col),
                Decision=decision if col == "id" else "RAW-ONLY",
                Notes="status and date_updated are NOT in this list — they import to availabilityStatus / updatedAt.",
            )
        )
    return out


def pending_rows(rows: list[dict]) -> list[dict]:
    ads = [r for r in rows if (r.get("Area_Details_Box") or "").strip()]
    ads_with_floor = 0
    for r in ads:
        if any(
            (r.get(f"Floor_{i}_Type") or r.get(f"Floor_{i}_Area") or r.get(f"Floor_{i}_Size") or "").strip()
            for i in range(1, 12)
        ):
            ads_with_floor += 1
    return [
        _row(
            Entity="PENDING BUSINESS DECISION",
            FinalField="(none)",
            FinalType="",
            ExistingOrNew="UNMAPPED",
            LegacySourceColumn="No_Rent_Price_Text",
            LegacySourceMeaning="HTML price note",
            MappingRule="Do not map into askingAmount or financialIncome",
            PopulatedLegacyRows=coverage(rows, "No_Rent_Price_Text"),
            ExampleLegacyValue=_ex(rows, "No_Rent_Price_Text"),
            Decision="PENDING",
            Notes="Possible later: notes / priceNote. Still on 01.",
        ),
        _row(
            Entity="PENDING BUSINESS DECISION",
            FinalField="(none)",
            FinalType="",
            ExistingOrNew="UNMAPPED",
            LegacySourceColumn="Deal",
            LegacySourceMeaning="Price basis (Negotiable / Asking / Net)",
            MappingRule="Do not map into askingAmount",
            PopulatedLegacyRows=coverage(rows, "Deal"),
            ExampleLegacyValue=_ex(rows, "Deal"),
            Decision="PENDING",
            Notes="Possible later: priceBasis. Still on 01.",
        ),
        _row(
            Entity="PENDING BUSINESS DECISION",
            FinalField="(none)",
            FinalType="",
            ExistingOrNew="UNMAPPED",
            LegacySourceColumn="Note",
            LegacySourceMeaning="HTML general notes",
            MappingRule="Do not stuff into JSON whitelist fields",
            PopulatedLegacyRows=coverage(rows, "Note"),
            ExampleLegacyValue=_ex(rows, "Note"),
            Decision="PENDING",
            Notes="Possible later: notes Text. Still on 01.",
        ),
        _row(
            Entity="PENDING BUSINESS DECISION",
            FinalField="(none)",
            FinalType="",
            ExistingOrNew="UNMAPPED",
            LegacySourceColumn="Area_Details_Box",
            LegacySourceMeaning="HTML hall/area narrative",
            MappingRule="NOT copied into units. Unique free-text (hall sizes, tables) beyond Floor_* slots.",
            PopulatedLegacyRows=coverage(rows, "Area_Details_Box"),
            ExampleLegacyValue=_ex(rows, "Area_Details_Box"),
            Decision="PENDING",
            Notes=f"{len(ads)} populated; only {ads_with_floor} also have Floor_1–11 slots. Remaining {len(ads)-ads_with_floor} have unique information not represented in units.",
        ),
    ]


def write_sheet(ws, rows: list[dict], surveys: list[dict]) -> dict:
    ws.sheet_properties.tabColor = "548235"
    ws["A1"] = "00_Final_Schema_Mapping — OUR FINAL FIELD  ←  OLD LEGACY SOURCE"
    ws["A1"].font = FONT_H
    ws.merge_cells("A1:L1")
    ws["A2"] = (
        "Green = NEW approved Survey field. Blue = EXISTING. Grey = NO LEGACY SOURCE. "
        "Purple = RAW-ONLY. Yellow = PENDING. "
        "Business_Developer → createdById, NOT SurveyContact.email. "
        "No permanent latitude/longitude Survey columns."
    )
    ws["A2"].font = FONT
    ws["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 36

    header_row = 3
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(header_row, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A4"
    ws.row_dimensions[header_row].height = 28

    sections = [
        ("A. SURVEY — 69 fields including id (12 NEW highlighted green)", survey_rows(rows, surveys)),
        ("B. SURVEYCONTACT — separate table; BD is not contact email", contact_rows(rows)),
        ("C. RAW / SYSTEM FIELDS — DO NOT IMPORT  (status and date_updated are NOT here)", raw_rows(rows)),
        ("D. PENDING BUSINESS DECISION — keep on 01_Legacy_Raw_Data", pending_rows(rows)),
    ]

    r = 4
    counts = {"survey": 0, "new": 0, "contact": 0, "raw": 0, "pending": 0}
    for title, block in sections:
        for c in range(1, 13):
            cell = ws.cell(r, c, title if c == 1 else None)
            cell.fill = SECTION_FILL
            cell.font = FONT_W
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        r += 1
        for item in block:
            for c, h in enumerate(HEADERS, 1):
                cell = ws.cell(r, c, item.get(h) if item.get(h) != "" else None)
                cell.font = FONT
                cell.alignment = WRAP
                kind = item.get("ExistingOrNew") or ""
                ent = item.get("Entity") or ""
                if item.get("FinalField") in APPROVED_NEW or kind == "NEW":
                    cell.fill = NEW_FILL
                    if c == 2:
                        cell.font = FONT_B
                elif "NO LEGACY" in str(item.get("LegacySourceColumn") or "") or "NO CONFIRMED" in str(item.get("LegacySourceColumn") or ""):
                    cell.fill = NONE_FILL
                elif "RAW" in ent or kind == "RAW-ONLY":
                    cell.fill = RAW_FILL
                elif "PENDING" in ent or item.get("Decision") == "PENDING":
                    cell.fill = PENDING_FILL
                elif kind.startswith("EXISTING"):
                    cell.fill = EXISTING_FILL
            if item.get("Entity") == "Survey":
                counts["survey"] += 1
                if item.get("FinalField") in APPROVED_NEW:
                    counts["new"] += 1
            elif item.get("Entity") == "SurveyContact":
                counts["contact"] += 1
            elif "RAW" in (item.get("Entity") or "") or item.get("Entity") == "MIGRATION METADATA ONLY":
                counts["raw"] += 1
            elif "PENDING" in (item.get("Entity") or ""):
                counts["pending"] += 1
            r += 1

    last = r - 1
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(12)}{last}"
    widths = [28, 26, 28, 22, 42, 28, 44, 14, 32, 28, 28, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return counts
