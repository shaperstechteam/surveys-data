#!/usr/bin/env python3
"""Generate the CRE backend seed workbook from the final clean workbook.

Usage (from the repo root):

    python migration/legacy-survey-migration/generate-cre-backend-seed.py

Reads (read-only):
    legacy-survey-final-clean-7398.xlsx     the validated historical/reference workbook

Writes:
    cre-backend-seed.xlsx
        01_Surveys          7,398 rows x 60 columns (legacySurveyId + 59 Survey fields)
        02_Survey_Contacts  copied unchanged
        03_Attachments      copied unchanged
        04_Seed_Exceptions  only the blockers/warnings that matter when seeding CRE

The seed Survey sheet drops askingAmount, possession, survey, advance, security,
gracePeriod, increment, agreementPeriod and surveyDate. Their historical values
stay in the final clean and audit workbooks. Every other value is copied exactly,
row for row by legacySurveyId, except that a blank zipCode receives the city-level
default postal code (CITY_ZIP) as text. Nothing is imported or uploaded.
"""
from __future__ import annotations

import sys

sys.dont_write_bytecode = True

from collections import Counter  # noqa: E402
from datetime import datetime  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "legacy-survey-final-clean-7398.xlsx"
OUTPUT = HERE / "cre-backend-seed.xlsx"

REMOVED_FIELDS = [
    "askingAmount", "possession", "survey", "advance", "security", "gracePeriod",
    "increment", "agreementPeriod", "surveyDate",
]
SEED_SURVEY_FIELDS = [
    "serialNumber", "propertyName", "address", "locationCoordinates", "city", "area", "zipCode",
    "yearBuilt", "yearRenovated", "buildingClass", "propertyType",
    "ownerName", "ownershipEntity", "propertyManager", "acquisitionDate",
    "plotArea", "areaUnit", "plotSize", "landShape", "topography", "zoning",
    "grossBuildingArea", "rentableArea", "usableArea", "coveredSize", "numberOfFloors", "ceilingHeight",
    "constructionType", "exteriorMaterial", "roofType", "foundationType", "parkingSpaces",
    "loadingDocks", "elevatorCount", "hvacDetails", "fireProtection",
    "availabilityStatus", "propertyStatus", "availableFor",
    "financialIncome", "financialExpenses", "financialMetrics", "occupancySurvey", "leaseSurvey",
    "utilityDetails", "units", "siteDetails", "buildingCondition", "environmentalDetails",
    "legalDetails", "marketDetails", "attachments",
    "createdById", "updatedById", "createdAt", "updatedAt", "freshSurveyCount",
    "neighboringBusinessIds", "submittedAt",
]
SURVEY_HEADERS = ["legacySurveyId"] + SEED_SURVEY_FIELDS

# City-level DEFAULT postal code (main GPO / main delivery office), not the exact
# neighbourhood code. Source: Pakistan Post "Post Code Directory of Delivery Post
# Offices" (pakpost.gov.pk/images/national post code directory.pdf); the office
# named here is the directory entry the code was taken from. Keys are matched after
# trimming and case-folding the city; the city text itself is never changed.
# Deliberately unmapped: Swat (a district, not one city), rawat and Mumtaz City
# (no matching directory entry).
CITY_ZIP = {
    "lahore": ("54000", "LAHORE GPO"),
    "islamabad": ("44000", "ISLAMABAD GPO"),
    "rawalpindi": ("46000", "RAWALPINDI GPO"),
    "karachi": ("74200", "KARACHI GPO"),
    "faisalabad": ("38000", "FAISALABAD GPO"),
    "peshawar": ("25000", "PESHAWAR GPO"),
    "jhelum": ("49600", "JHELUM GPO"),
    "gujranwala": ("52250", "GUJRANWALA GPO"),
    "gujrat": ("50700", "GUJRAT GPO"),
    "sialkot": ("51310", "SIALKOT GPO"),
    "multan": ("60000", "MULTAN GPO"),
    "mirpur ajk": ("10250", "MIRPUR GPO (Azad Kashmir)"),
    "abbottabad": ("22010", "ABBOTTABAD GPO"),
    "sargodha": ("40100", "SARGODHA GPO"),
    "wah cantt.": ("47040", "WAH CANTT. GPO"),
    "mardan": ("23200", "MARDAN GPO"),
    "hyderabad": ("71000", "HYDERABAD GPO"),
    "mandi bahauddin": ("50400", "MANDI BAHAUDDIN GPO"),
    "kasur": ("55050", "KASUR GPO"),
    "rahim yar khan": ("64200", "RAHIMYAR KHAN GPO"),
    "sahiwal": ("57000", "SAHIWAL GPO"),
    "sukkur": ("65200", "SUKKUR GPO"),
    "murree": ("47150", "MURREE GPO"),
    "attock": ("43600", "ATTOCK GPO"),
    "sheikhupura": ("39350", "QILA SHEIKHUPURA GPO"),
    "okara": ("56300", "OKARA GPO"),
    "haripur": ("22620", "HARIPUR GPO"),
    "bahawalpur": ("63100", "BAHAWALPUR GPO"),
    "quetta": ("87300", "QUETTA GPO"),
    "muzaffarabad": ("13100", "MUZAFFARABAD GPO"),
    "dera ghazi khan": ("32200", "DERA GHAZI KHAN GPO"),
    "daska": ("51010", "DASKA"),
    "narowal": ("51600", "NAROWAL GPO"),
    "vehari": ("61100", "VEHARI GPO"),
    "burewala": ("61010", "BUREWALA"),
    "chakwal": ("48800", "CHAKWAL GPO"),
    "hafizabad": ("52110", "HAFIZ ABAD"),
    "wazirabad": ("52000", "WAZIRABAD"),
    "jhang": ("35200", "JHANG GPO"),
    "taxila": ("47080", "TAXILA"),
    "wah": ("47000", "WAH"),
    "mianwali": ("42200", "MIANWALI GPO"),
    "fateh jang": ("43350", "FATEH JANG"),
    "muridke": ("39000", "MURIDKE"),
    "kharian": ("50090", "KHARIAN CITY"),
    "lalamusa": ("50200", "LALA MUSA"),
    "bahawalnagar": ("62300", "BAHAWAL NAGAR GPO"),
    "lodhran": ("59320", "LODHRAN"),
    "toba tek singh": ("36050", "TOBA TAKE SINGH GPO"),
    "khanewal": ("58150", "KHANEWAL GPO"),
    "gujar khan": ("47850", "GUJAR KHAN GPO"),
    "mirpur khas": ("69000", "MIRPUR KHAS GPO"),
    "dina": ("49400", "DINA"),
    "mansehra": ("21300", "MANSEHRA GPO"),
    "layyah": ("31200", "LAYYAH GPO"),
    "sarai alamgir": ("50000", "SARAI ALAMGIR"),
    "chiniot": ("35400", "CHINIOT"),
    "pir mahal": ("36300", "PIR MAHAL"),
    "dara adam khel": ("26100", "DARA ADAM KHEL"),
    "hazro": ("43440", "HAZRO"),
    "arifwala": ("57450", "ARIF WALA"),
    "dinga": ("50280", "DINGA"),
    "bannu": ("28100", "BANNU GPO"),
    "phalia": ("50430", "PHALIA"),
    "joharabad": ("41200", "JAUHAR ABAD"),
    "hattar": ("43394", "HATTAR"),
    "nowshera": ("24100", "NOWSHERA GPO"),
    "ali pur chatha": ("52080", "ALI PUR CHATHA"),
}


def city_zip(city):
    if city is None or not str(city).strip():
        return None
    hit = CITY_ZIP.get(str(city).strip().casefold())
    return hit[0] if hit else None
EXCEPTION_HEADERS = ["entity", "legacySurveyId", "field", "rawValue", "normalizedValue",
                     "reason", "severity", "notes"]

# per-row exceptions that still matter for a seed without the removed fields
KEEP_ROW_REASONS = {
    "INVALID_LEGACY_DATE", "BUSINESS_DEVELOPER_MISSING", "MISSING_COORDINATES", "INVALID_COORDINATES",
    "COORDINATES_OUTSIDE_PAKISTAN", "CONFLICT_WITH_POSSESSION", "DUPLICATE_CANDIDATE",
    "UNIT_FLOOR_LABEL_MISSING", "UNPARSEABLE_FLOOR_AREA",
    "REVIEW_CP_SEMANTICS", "INVALID_NAME", "REVIEW_MULTI_PHONE",          # former HOLD contacts
}
KEEP_DATASET_REASONS = {
    "CREATED_BY_ID_REQUIRES_USER_LOOKUP", "TIMESTAMP_TIMEZONE_UNCONFIRMED", "LEGACY_STATUS_CONSTANT_ACTIVE",
}


def read_sheet(wb, name):
    it = wb[name].iter_rows(values_only=True)
    headers = list(next(it))
    return headers, [dict(zip(headers, r)) for r in it if any(v is not None for v in r)]


def build():
    src = load_workbook(SOURCE, read_only=True)
    s_head, surveys = read_sheet(src, "01_Surveys")
    c_head, contacts = read_sheet(src, "02_Survey_Contacts")
    a_head, attachments = read_sheet(src, "03_Attachments")
    _, exceptions = read_sheet(src, "04_Exceptions")

    missing = [f for f in SURVEY_HEADERS if f not in s_head]
    if missing:
        raise SystemExit(f"final clean workbook lacks {missing}")
    if len(surveys) != 7398:
        raise SystemExit(f"expected 7,398 surveys, found {len(surveys)}")
    seed_surveys = [{h: r[h] for h in SURVEY_HEADERS} for r in surveys]
    # zipCode: keep any existing value; otherwise the city-level default (as text)
    for r in seed_surveys:
        if r["zipCode"] in (None, ""):
            r["zipCode"] = city_zip(r["city"])

    # ------------------------------------------------ seed exceptions
    seed_exc = []
    for e in exceptions:
        if e["entity"] == "Dataset" and e["reason"] in KEEP_DATASET_REASONS:
            seed_exc.append(dict(e))
        elif e["entity"] in ("Survey", "SurveyContact") and e["reason"] in KEEP_ROW_REASONS:
            row = dict(e)
            if row["field"] == "surveyDate/createdAt":
                row["field"] = "createdAt"
                row["notes"] = "No valid date_created; createdAt is blank in the seed"
            seed_exc.append(row)

    reasons = Counter(e["reason"] for e in exceptions)
    blank_type = sum(1 for r in seed_surveys if r["propertyType"] in (None, ""))
    att_excluded = sum(1 for e in exceptions if e["entity"] == "Attachment")

    def dataset(field, raw, reason, severity, notes):
        seed_exc.append({"entity": "Dataset", "legacySurveyId": None, "field": field, "rawValue": raw,
                         "normalizedValue": None, "reason": reason, "severity": severity, "notes": notes})

    dataset("surveyDate", None, "SCHEMA_REMOVE_SURVEY_DATE", "HIGH",
            "surveyDate is not in this seed. CRE Survey.surveyDate (currently required) must be removed "
            "from the Prisma schema/backend before import; createdAt carries the historical timestamp")
    dataset("area, plotSize, coveredSize, availableFor", None, "SCHEMA_ADD_SURVEY_FIELDS", "HIGH",
            "These seed columns must exist on the CRE Survey model before import")
    dataset("locationCoordinates", None, "SCHEMA_LOCATION_COORDINATES_JSON", "HIGH",
            'Values are JSON {"latitude":…,"longitude":…}; CRE column must accept Json, not Decimal(9,6)')
    dataset("units", None, "SCHEMA_UNITS_DIMENSIONS", "MEDIUM",
            "units JSON items may carry dimensions; extend SurveyUnitDto or the key will be stripped")
    dataset("createdAt", str(reasons["INVALID_LEGACY_DATE"]), "CREATED_AT_MISSING", "HIGH",
            f"{reasons['INVALID_LEGACY_DATE']:,} surveys have no valid legacy timestamp (rows below); "
            "decide: import with createdAt = import time, or hold them")
    dataset("propertyType", str(blank_type), "PROPERTY_TYPE_UNRESOLVED", "MEDIUM",
            f"{blank_type:,} surveys have blank propertyType (legacy 'building' or blank has no safe enum)")
    dataset("phone", str(reasons["CONTACT_PHONE_INCOMPLETE_OR_INVALID"]), "CONTACT_PHONE_RAW", "MEDIUM",
            f"Contact phones are raw legacy text; {reasons['CONTACT_PHONE_INCOMPLETE_OR_INVALID']:,} are "
            f"incomplete/invalid and {reasons['CONTACT_MULTIPLE_PHONES']:,} hold several numbers. "
            "Importer must normalise or accept as-is")
    unmapped = Counter(str(r["city"]).strip() for r in seed_surveys
                       if r["zipCode"] is None and r["city"] not in (None, ""))
    no_city = sum(1 for r in seed_surveys if r["city"] in (None, ""))
    dataset("zipCode", None, "ZIPCODE_CITY_LEVEL_DEFAULT", "LOW",
            "zipCode is the city's main GPO/delivery-office code from the Pakistan Post directory, not the "
            f"exact neighbourhood code. Left blank: {no_city} surveys with no city; unmapped cities "
            + ", ".join(f"{c} ({n})" for c, n in unmapped.most_common()))
    dataset("attachments", str(att_excluded), "ATTACHMENTS_NOT_IN_MANIFEST", "MEDIUM",
            f"{att_excluded:,} legacy file references are not in 03_Attachments (placeholder, unverified, "
            "unsupported or unavailable). Manifest files still need upload to S3 before AttachmentRecords exist")

    order = {"Dataset": 0, "Survey": 1, "SurveyContact": 2}
    seed_exc.sort(key=lambda e: (order[e["entity"]], e["legacySurveyId"] or 0, e["field"] or "", e["reason"]))

    write(seed_surveys, c_head, contacts, a_head, attachments, seed_exc)
    print(f"Wrote {OUTPUT.name}: surveys={len(seed_surveys)} x {len(SURVEY_HEADERS)} cols, "
          f"contacts={len(contacts)}, attachments={len(attachments)}, seedExceptions={len(seed_exc)}")
    print("seed exceptions by reason:", dict(Counter(e["reason"] for e in seed_exc).most_common()))


# ---------------------------------------------------------------- write

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOP = Alignment(vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)


def write_sheet(wb, title, headers, records, widths=None, wrap=()):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    wrap = set(wrap)
    for r, rec in enumerate(records, 2):
        for c, h in enumerate(headers, 1):
            val = rec.get(h)
            if val == "":
                val = None
            cell = ws.cell(r, c)
            cell.value = val
            if isinstance(val, str):
                cell.data_type = "s"
            elif isinstance(val, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cell.alignment = WRAP if h in wrap else TOP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(h, max(12, min(len(h) + 4, 24)))
    ws.sheet_view.zoomScale = 90


def write(surveys, c_head, contacts, a_head, attachments, seed_exc):
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "01_Surveys", SURVEY_HEADERS, surveys,
                widths={"address": 32, "locationCoordinates": 34, "area": 28, "ownerName": 26,
                        "units": 48, "marketDetails": 48, "createdById": 30, "createdAt": 20},
                wrap={"address", "locationCoordinates", "units", "marketDetails"})
    write_sheet(wb, "02_Survey_Contacts", c_head, contacts,
                widths={"name": 28, "role": 24, "phone": 28, "createdAt": 20})
    write_sheet(wb, "03_Attachments", a_head, attachments,
                widths={"originalUrl": 70, "originalFilename": 36, "mimeType": 16})
    write_sheet(wb, "04_Seed_Exceptions", EXCEPTION_HEADERS, seed_exc,
                widths={"entity": 14, "field": 26, "rawValue": 50, "normalizedValue": 30,
                        "reason": 34, "severity": 10, "notes": 60},
                wrap={"rawValue", "normalizedValue", "notes"})
    wb.save(OUTPUT)


if __name__ == "__main__":
    build()
