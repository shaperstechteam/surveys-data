#!/usr/bin/env python3
"""Reopen the generated workbook and check it against the source CSV."""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source" / "merged_properties_final.csv"
CANDIDATES = [
    HERE / "legacy-survey-migration-7398.xlsx",
    HERE / "legacy-survey-migration-7398-approved.xlsx",
]
XLSX = next((p for p in CANDIDATES if p.exists()), HERE / "legacy-survey-migration-7398.xlsx")
EXPECTED = 7398
REQUIRED_SHEETS = [
    "00_Final_Schema_Mapping",
    "01_Legacy_Raw_Data",
    "02_Survey_Normalized",
    "03_Survey_Contacts",
    "04_Attachments",
    "05_Field_Mapping",
    "06_Enum_Mappings",
    "07_JSON_Mappings",
    "08_Unmapped_Fields",
    "09_Data_Issues",
    "10_Duplicates",
    "11_Decision_Register",
    "12_Migration_Summary",
]


def data_rows(ws, header_row=1):
    n = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(r, 1).value is not None:
            n += 1
    return n


def col_index(ws, name, header_row=1):
    for c in range(1, ws.max_column + 1):
        if ws.cell(header_row, c).value == name:
            return c
    raise SystemExit(f"Missing column {name} on {ws.title}")


def main():
    if not XLSX.exists():
        raise SystemExit(f"Workbook not found: {XLSX}")
    with SOURCE.open(encoding="utf-8-sig", newline="") as fh:
        src = list(csv.DictReader(fh))
    src_ids = {int(r["id"]) for r in src}
    src_keys = list(src[0].keys())
    src_by_id = {int(r["id"]): r for r in src}

    print(f"Loading {XLSX}")
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    errors = []
    notes = []

    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        errors.append(f"Missing sheets: {missing}")
    extra = [s for s in wb.sheetnames if s not in REQUIRED_SHEETS]
    if extra:
        notes.append(f"Extra sheets: {extra}")

    raw = wb["01_Legacy_Raw_Data"]
    raw_headers = [c.value for c in next(raw.iter_rows(min_row=1, max_row=1))]
    if raw_headers != src_keys:
        lost = [k for k in src_keys if k not in raw_headers]
        extra_h = [k for k in raw_headers if k not in src_keys]
        if lost:
            errors.append(f"Raw sheet lost columns: {lost}")
        if extra_h:
            notes.append(f"Raw sheet extra columns: {extra_h}")
    raw_ids = []
    raw_by_id = {}
    for row in raw.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        raw_ids.append(int(row[0]))
        raw_by_id[int(row[0])] = row
    if len(raw_ids) != EXPECTED:
        errors.append(f"01_Legacy_Raw_Data rows={len(raw_ids)} expected {EXPECTED}")
    if len(set(raw_ids)) != len(raw_ids):
        errors.append("01_Legacy_Raw_Data has duplicate ids")
    if set(raw_ids) != src_ids:
        errors.append(f"Raw ids differ from source: missing {len(src_ids - set(raw_ids))} extra {len(set(raw_ids) - src_ids)}")

    norm = wb["02_Survey_Normalized"]
    norm_header = [c.value for c in next(norm.iter_rows(min_row=1, max_row=1))]
    for forbidden in ("latitude", "longitude", "legacyPropertyType", "listingIntent", "locality"):
        if forbidden in norm_header:
            errors.append(f"02 contains forbidden destination column: {forbidden}")
    for required in (
        "legacySurveyId", "area", "plotSize", "coveredSize", "availableFor", "askingAmount",
        "possession", "survey", "advance", "security", "gracePeriod", "increment",
        "agreementPeriod", "availabilityStatus", "propertyStatus", "updatedAt",
        "createdById", "userMatchStatus", "businessDeveloperEmail", "locationCoordinates",
        "city", "units",
    ):
        if required not in norm_header:
            errors.append(f"02 missing required column: {required}")
    id_i = norm_header.index("legacySurveyId")
    status_i = norm_header.index("migrationStatus")
    coord_i = norm_header.index("locationCoordinates")
    units_i = norm_header.index("units")
    market_i = norm_header.index("marketDetails")
    serial_i = norm_header.index("serialNumber")
    map_sheet = wb["05_Field_Mapping"]
    map_headers = [c.value for c in next(map_sheet.iter_rows(min_row=1, max_row=1))]
    lf_i = map_headers.index("legacyField")
    mapped_cols = []
    for row in map_sheet.iter_rows(min_row=2, values_only=True):
        if row[lf_i]:
            mapped_cols.append(row[lf_i])
    missing_map = [k for k in src_keys if k not in mapped_cols]
    if missing_map:
        errors.append(f"05_Field_Mapping missing columns: {missing_map}")
    norm_ids = []
    statuses = {}
    formula_hits = 0
    json_fail = 0
    for row in norm.iter_rows(min_row=2, values_only=True):
        if row[id_i] is None:
            continue
        lid = int(row[id_i])
        norm_ids.append(lid)
        statuses[lid] = row[status_i]
        for cell in row:
            if isinstance(cell, str) and cell.startswith("="):
                formula_hits += 1
        for idx in (coord_i, units_i, market_i):
            val = row[idx]
            if val:
                try:
                    json.loads(val)
                except Exception:
                    json_fail += 1
        if row[serial_i] not in (None, ""):
            notes.append(f"serialNumber pre-filled for {lid}")
    if len(norm_ids) != EXPECTED:
        errors.append(f"02_Survey_Normalized rows={len(norm_ids)} expected {EXPECTED}")
    if set(norm_ids) != src_ids:
        errors.append("Normalized ids do not match source")
    if formula_hits:
        errors.append(f"Unexpected formulas on 02: {formula_hits}")
    if json_fail:
        errors.append(f"Invalid JSON cells on 02: {json_fail}")

    contacts = wb["03_Survey_Contacts"]
    ch = [c.value for c in next(contacts.iter_rows(min_row=1, max_row=1))]
    cid = ch.index("legacySurveyId")
    contact_ids = []
    for row in contacts.iter_rows(min_row=2, values_only=True):
        if row[cid] is None:
            continue
        contact_ids.append(int(row[cid]))
    orphan_c = set(contact_ids) - src_ids
    if orphan_c:
        errors.append(f"Contacts reference unknown surveys: {len(orphan_c)}")

    atts = wb["04_Attachments"]
    ah = [c.value for c in next(atts.iter_rows(min_row=1, max_row=1))]
    aid = ah.index("legacySurveyId")
    att_n = 0
    orphan_a = 0
    for row in atts.iter_rows(min_row=2, values_only=True):
        if row[aid] is None:
            continue
        att_n += 1
        if int(row[aid]) not in src_ids:
            orphan_a += 1
    if orphan_a:
        errors.append(f"Attachments reference unknown surveys: {orphan_a}")

    # Sample compare first / mid / last / floors / attachments / bad date / multi-phone
    sample_ids = [int(src[0]["id"]), int(src[len(src) // 2]["id"]), int(src[-1]["id"])]
    for r in src:
        if any((r.get(f"Floor_{i}_Type") or "").strip() for i in range(4, 12)):
            sample_ids.append(int(r["id"]))
            break
    for r in src:
        pics = r.get("Pictures") or ""
        if pics.startswith("[") and pics.count("http") >= 3:
            sample_ids.append(int(r["id"]))
            break
    for r in src:
        if (r.get("date_created") or "").startswith("0000-00-00"):
            sample_ids.append(int(r["id"]))
            break
    for r in src:
        cn = r.get("Contact_Number") or ""
        if "/" in cn or cn.count("03") >= 2:
            sample_ids.append(int(r["id"]))
            break

    raw_h = raw_headers
    owner_i = raw_h.index("Owner_Name")
    for sid in sample_ids:
        src_row = src_by_id[sid]
        excel_row = raw_by_id.get(sid)
        if not excel_row:
            errors.append(f"Sample id {sid} missing from raw sheet")
            continue
        if str(excel_row[owner_i] or "") != (src_row.get("Owner_Name") or ""):
            errors.append(f"Owner_Name mismatch for {sid}")
        if statuses.get(sid) is None:
            errors.append(f"Sample id {sid} missing from normalized sheet")

    if wb.sheetnames[0] != "00_Final_Schema_Mapping":
        errors.append(f"First sheet is {wb.sheetnames[0]}, expected 00_Final_Schema_Mapping")
    s0 = wb["00_Final_Schema_Mapping"]
    new12 = {
        "area", "plotSize", "coveredSize", "availableFor", "askingAmount",
        "possession", "survey", "advance", "security", "gracePeriod",
        "increment", "agreementPeriod",
    }
    found_new = set()
    survey_fields = set()
    contact_fields = set()
    raw_fields = 0
    pending_fields = 0
    for row in s0.iter_rows(min_row=4, values_only=True):
        ent, field = row[0], row[1]
        if ent == "Survey" and field:
            survey_fields.add(field)
            if field in new12:
                found_new.add(field)
        elif ent == "SurveyContact" and field:
            contact_fields.add(field)
        elif ent in ("RAW / SYSTEM — DO NOT IMPORT", "MIGRATION METADATA ONLY"):
            raw_fields += 1
        elif ent == "PENDING BUSINESS DECISION":
            pending_fields += 1
    missing_new = sorted(new12 - found_new)
    if missing_new:
        errors.append(f"00 missing approved NEW fields: {missing_new}")
    if "id" not in survey_fields:
        errors.append("00 Survey section missing id")

    report = {
        "ok": not errors,
        "errors": errors,
        "notes": notes,
        "rawRows": len(raw_ids),
        "normalizedRows": len(norm_ids),
        "contactRows": len(contact_ids),
        "attachmentRows": att_n,
        "sampleIdsChecked": sample_ids,
        "sheet00": {
            "firstSheet": wb.sheetnames[0],
            "surveyFields": len(survey_fields),
            "approvedNewShown": sorted(found_new),
            "surveyContactFields": sorted(contact_fields),
            "rawOrMetadataRows": raw_fields,
            "pendingRows": pending_fields,
        },
        "workbook": str(XLSX),
    }
    out = HERE / "validation-report.json"
    out.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
